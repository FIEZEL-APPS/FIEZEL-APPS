"""Question Engine — Question DNA, ingestion, validasi, versioning, varian, lifecycle.

Lifecycle: DRAFT -> REVIEW -> APPROVED -> PUBLISHED -> ARCHIVED.
Hasil parser/import/kandidat otomatis selalu masuk REVIEW; hanya guru yang menerbitkan.
"""
import io
import csv
import re
import uuid
import hashlib
from datetime import datetime, timezone
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from pydantic import BaseModel, Field

from db import db
from auth import teacher_user
from curriculum import prerequisites_of

router = APIRouter(prefix="/api/questions", tags=["questions"])

STATUSES = ["DRAFT", "REVIEW", "APPROVED", "PUBLISHED", "ARCHIVED"]
COGNITIVE = ["C1", "C2", "C3", "C4", "C5", "C6"]
QTYPES = ["mcq", "true_false", "short_answer", "numeric", "essay", "cloze"]


def now():
    return datetime.now(timezone.utc)


def norm_stem(s: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", "", (s or "").lower())).strip()


def stem_hash(s: str) -> str:
    return hashlib.sha1(norm_stem(s).encode()).hexdigest()[:16]


class QuestionIn(BaseModel):
    competency_id: str
    stem: str
    question_type: str = "mcq"
    options: list[str] = Field(default_factory=list)
    answer_key: str = ""
    explanation: str = ""
    hints: list[str] = Field(default_factory=list)
    difficulty: int = 2
    cognitive_level: str = "C2"
    skill: str = ""
    misconception_id: str | None = None
    distractor_misconceptions: dict[str, str] = Field(default_factory=dict)
    prerequisite_competency_ids: list[str] = Field(default_factory=list)
    estimated_time: int = 60
    is_transfer: bool = False
    topic_id: str | None = None
    material_id: str | None = None
    variant_group_id: str | None = None
    source: str = "manual"
    status: str = "DRAFT"
    tags: list[str] = Field(default_factory=list)


async def _dna(body: QuestionIn, actor: str, question_id: str | None = None, version: int = 1) -> dict:
    comp = await db.curriculum_nodes.find_one({"id": body.competency_id, "type": "competency"}, {"_id": 0})
    if not comp:
        raise HTTPException(404, f"competency {body.competency_id} tidak ditemukan")
    prereq = body.prerequisite_competency_ids or await prerequisites_of(body.competency_id)
    qid = question_id or f"Q-{uuid.uuid4().hex[:10].upper()}"
    doc = {
        "question_id": qid, "version": version, "is_current": True,
        "status": body.status if body.status in STATUSES else "DRAFT",
        "curriculum_id": comp.get("curriculum_id"), "phase_id": comp.get("phase_id"),
        "grade_id": comp.get("grade_id"), "subject_id": comp.get("subject_id"),
        "element_id": comp.get("element_id"), "cp_id": comp.get("cp_id"),
        "tp_id": comp.get("tp_id"), "indicator_id": comp.get("indicator_id"),
        "competency_id": comp["id"], "competency_name": comp["name"],
        "topic_id": body.topic_id, "material_id": body.material_id,
        "skill": body.skill or (comp.get("meta") or {}).get("legacy_skill") or comp["code"],
        "difficulty": max(1, min(5, int(body.difficulty))),
        "cognitive_level": body.cognitive_level if body.cognitive_level in COGNITIVE else "C2",
        "question_type": body.question_type if body.question_type in QTYPES else "mcq",
        "stem": body.stem.strip(), "stem_hash": stem_hash(body.stem),
        "options": [o.strip() for o in body.options if str(o).strip()],
        "answer_key": str(body.answer_key).strip(), "explanation": body.explanation.strip(),
        "hints": [h for h in body.hints if h],
        "misconception_id": body.misconception_id,
        "distractor_misconceptions": body.distractor_misconceptions,
        "prerequisite_competency_ids": prereq,
        "estimated_time": max(10, int(body.estimated_time)),
        "is_transfer": bool(body.is_transfer),
        "variant_group_id": body.variant_group_id or qid,
        "source": body.source, "tags": body.tags,
        "created_by": actor, "created_at": now(), "updated_at": now(),
    }
    return doc


# ------------------------- validasi -------------------------
async def validate_doc(doc: dict) -> list[dict]:
    issues: list[dict] = []

    def add(level, code, msg):
        issues.append({"level": level, "code": code, "message": msg})

    if len(doc.get("stem", "")) < 8:
        add("error", "stem_too_short", "Pertanyaan terlalu pendek / tidak lengkap.")
    if doc["question_type"] == "mcq":
        if len(doc.get("options") or []) < 2:
            add("error", "options_missing", "Soal pilihan ganda butuh minimal 2 opsi.")
        keys = [chr(65 + i) for i in range(len(doc.get("options") or []))]
        if doc.get("answer_key", "").upper() not in keys:
            add("error", "answer_key_invalid",
                f"Kunci jawaban “{doc.get('answer_key')}” tidak cocok dengan opsi ({', '.join(keys)}).")
        opts = [norm_stem(o) for o in doc.get("options") or []]
        if len(set(opts)) != len(opts):
            add("warn", "duplicate_options", "Ada opsi jawaban yang identik.")
    elif not doc.get("answer_key"):
        add("error", "answer_key_missing", "Kunci jawaban belum diisi.")
    if not doc.get("competency_id"):
        add("error", "competency_missing", "Soal ini belum terhubung ke kompetensi.")
    if not doc.get("tp_id"):
        add("error", "tp_missing", "Soal ini belum terhubung ke TP.")
    if not doc.get("cp_id"):
        add("warn", "cp_missing", "Soal ini belum terhubung ke CP.")
    if not doc.get("indicator_id"):
        add("warn", "indicator_missing", "Kompetensi soal ini belum berada di bawah indikator.")
    if not doc.get("explanation"):
        add("warn", "explanation_missing", "Belum ada penjelasan — murid yang salah tidak akan mendapat pembahasan.")
    if not doc.get("hints"):
        add("info", "hint_missing", "Belum ada petunjuk (hint) untuk scaffolding saat murid gagal.")
    if doc["question_type"] == "mcq" and not doc.get("distractor_misconceptions") and not doc.get("misconception_id"):
        add("info", "misconception_missing",
            "Belum ada pemetaan miskonsepsi — diagnosis kesalahan jadi kurang tajam.")
    if not doc.get("prerequisite_competency_ids"):
        add("info", "prerequisite_missing", "Kompetensi ini belum punya prasyarat; remediasi akan kurang terarah.")
    if doc["difficulty"] < 1 or doc["difficulty"] > 5:
        add("error", "difficulty_invalid", "Tingkat kesulitan harus 1–5.")
    if doc["cognitive_level"] not in COGNITIVE:
        add("error", "cognitive_invalid", "Level kognitif harus C1–C6.")
    dup = await db.questions.find_one({"stem_hash": doc["stem_hash"], "is_current": True,
                                        "question_id": {"$ne": doc["question_id"]},
                                        "status": {"$ne": "ARCHIVED"}}, {"_id": 0, "question_id": 1})
    if dup:
        add("warn", "duplicate_question", f"Soal ini mirip/duplikat dengan {dup['question_id']}.")
    same = await db.questions.count_documents({"competency_id": doc["competency_id"], "is_current": True,
                                               "question_id": {"$ne": doc["question_id"]},
                                               "status": {"$in": ["APPROVED", "PUBLISHED"]}})
    if same >= 4:
        add("info", "competency_saturated",
            f"Soal ini mengukur kompetensi yang sama dengan {same} soal lain — pertimbangkan variasi kognitif/transfer.")
    return issues


def blocking(issues: list[dict]) -> bool:
    return any(i["level"] == "error" for i in issues)


@router.post("/validate")
async def validate_preview(body: QuestionIn, u=Depends(teacher_user)):
    doc = await _dna(body, u["user_id"])
    issues = await validate_doc(doc)
    return {"issues": issues, "publishable": not blocking(issues)}


# ------------------------- CRUD + versioning -------------------------
@router.post("")
async def create_question(body: QuestionIn, u=Depends(teacher_user)):
    doc = await _dna(body, u["user_id"])
    issues = await validate_doc(doc)
    doc["issues"] = issues
    if doc["status"] in ("APPROVED", "PUBLISHED") and blocking(issues):
        doc["status"] = "REVIEW"
    await db.questions.insert_one(dict(doc))
    doc.pop("_id", None)
    return {"question": doc, "issues": issues}


@router.get("")
async def list_questions(tp_id: str | None = None, competency_id: str | None = None,
                         status: str | None = None, subject_id: str | None = None,
                         grade_id: str | None = None, is_transfer: bool | None = None,
                         q: str | None = None, limit: int = 200, all_versions: bool = False):
    query: dict[str, Any] = {}
    if not all_versions:
        query["is_current"] = True
    for k, v in (("tp_id", tp_id), ("competency_id", competency_id), ("status", status),
                 ("subject_id", subject_id), ("grade_id", grade_id)):
        if v:
            query[k] = v
    if is_transfer is not None:
        query["is_transfer"] = is_transfer
    if q:
        query["stem"] = {"$regex": re.escape(q), "$options": "i"}
    return await db.questions.find(query, {"_id": 0}).sort("created_at", -1).to_list(min(limit, 1000))


@router.get("/{question_id}")
async def get_question(question_id: str):
    doc = await db.questions.find_one({"question_id": question_id, "is_current": True}, {"_id": 0})
    if not doc:
        raise HTTPException(404, "soal tidak ditemukan")
    return doc


@router.get("/{question_id}/history")
async def question_history(question_id: str):
    return await db.questions.find({"question_id": question_id}, {"_id": 0}).sort("version", 1).to_list(100)


@router.put("/{question_id}")
async def update_question(question_id: str, body: QuestionIn, u=Depends(teacher_user)):
    """Edit = versi baru. Versi lama disimpan agar evidence murid tidak rusak."""
    cur = await db.questions.find_one({"question_id": question_id, "is_current": True}, {"_id": 0})
    if not cur:
        raise HTTPException(404, "soal tidak ditemukan")
    doc = await _dna(body, u["user_id"], question_id=question_id, version=cur["version"] + 1)
    doc["created_at"] = cur["created_at"]
    doc["variant_group_id"] = body.variant_group_id or cur.get("variant_group_id") or question_id
    issues = await validate_doc(doc)
    doc["issues"] = issues
    if doc["status"] in ("APPROVED", "PUBLISHED") and blocking(issues):
        doc["status"] = "REVIEW"
    await db.question_versions.insert_one({**{k: v for k, v in cur.items()}, "archived_at": now()})
    await db.questions.update_one({"question_id": question_id, "version": cur["version"]},
                                  {"$set": {"is_current": False}})
    await db.questions.insert_one(dict(doc))
    doc.pop("_id", None)
    return {"question": doc, "issues": issues, "previous_version": cur["version"]}


class ReviewIn(BaseModel):
    action: str  # approve | publish | request_changes | archive | draft
    note: str = ""


@router.post("/{question_id}/review")
async def review_question(question_id: str, body: ReviewIn, u=Depends(teacher_user)):
    doc = await get_question(question_id)
    target = {"approve": "APPROVED", "publish": "PUBLISHED", "request_changes": "REVIEW",
              "archive": "ARCHIVED", "draft": "DRAFT"}.get(body.action)
    if not target:
        raise HTTPException(400, "action tidak dikenal")
    if target in ("APPROVED", "PUBLISHED"):
        issues = await validate_doc(doc)
        if blocking(issues):
            raise HTTPException(400, {"message": "Soal belum lolos validasi", "issues": issues})
    await db.questions.update_one({"question_id": question_id, "is_current": True},
                                  {"$set": {"status": target, "updated_at": now(),
                                            "reviewed_by": u["user_id"], "review_note": body.note}})
    return await get_question(question_id)


class BulkReviewIn(BaseModel):
    question_ids: list[str]
    action: str


@router.post("/bulk-review")
async def bulk_review(body: BulkReviewIn, u=Depends(teacher_user)):
    out = {"ok": [], "failed": []}
    for qid in body.question_ids:
        try:
            await review_question(qid, ReviewIn(action=body.action), u)
            out["ok"].append(qid)
        except HTTPException as e:
            out["failed"].append({"question_id": qid, "detail": e.detail})
    return out


# ------------------------- varian -------------------------
NUM_RE = re.compile(r"-?\d+(?:[.,]\d+)?")


class VariantIn(BaseModel):
    count: int = 2
    transfer: bool = False


@router.post("/{question_id}/variants")
async def make_variants(question_id: str, body: VariantIn, u=Depends(teacher_user)):
    """Varian deterministik untuk kompetensi yang sama (anti hafalan, mendukung retry & transfer)."""
    src = await get_question(question_id)
    made = []
    for i in range(max(1, min(6, body.count))):
        stem = src["stem"]
        shift = i + 1
        if src["question_type"] in ("numeric", "short_answer") and NUM_RE.search(stem):
            stem = NUM_RE.sub(lambda m: str(_shift_num(m.group(0), shift)), stem)
        prefix = "Dalam situasi berbeda: " if body.transfer else f"Variasi {shift}: "
        cand = QuestionIn(
            competency_id=src["competency_id"], stem=prefix + stem,
            question_type=src["question_type"], options=list(src.get("options") or []),
            answer_key=src.get("answer_key", ""), explanation=src.get("explanation", ""),
            hints=list(src.get("hints") or []), difficulty=min(5, src["difficulty"] + (1 if body.transfer else 0)),
            cognitive_level="C3" if body.transfer and src["cognitive_level"] in ("C1", "C2") else src["cognitive_level"],
            skill=src.get("skill", ""), misconception_id=src.get("misconception_id"),
            distractor_misconceptions=src.get("distractor_misconceptions") or {},
            prerequisite_competency_ids=src.get("prerequisite_competency_ids") or [],
            estimated_time=src["estimated_time"], is_transfer=body.transfer,
            topic_id=src.get("topic_id"), material_id=src.get("material_id"),
            variant_group_id=src.get("variant_group_id") or question_id,
            source="variant_generator", status="REVIEW")
        doc = await _dna(cand, u["user_id"])
        doc["issues"] = await validate_doc(doc)
        doc["derived_from"] = question_id
        await db.questions.insert_one(dict(doc))
        doc.pop("_id", None)
        made.append(doc)
    return {"created": made, "note": "Varian masuk status REVIEW — guru menerbitkan."}


def _shift_num(raw: str, shift: int):
    v = raw.replace(",", ".")
    try:
        if "." in v:
            return round(float(v) + shift, 2)
        return int(v) + shift
    except ValueError:
        return raw


# ------------------------- ingestion -------------------------
PASTE_SPLIT = re.compile(r"\n\s*(?=\d{1,3}\s*[.)]\s)")
OPT_RE = re.compile(r"^\s*([A-Ea-e])\s*[.)]\s*(.+)$")
KEY_RE = re.compile(r"^\s*(?:jawaban|kunci|answer|key)\s*[:\-]?\s*([A-Ea-e0-9].*)$", re.I)
EXP_RE = re.compile(r"^\s*(?:penjelasan|pembahasan|explanation|alasan)\s*[:\-]?\s*(.*)$", re.I)
HINT_RE = re.compile(r"^\s*(?:hint|petunjuk)\s*[:\-]?\s*(.*)$", re.I)
MIS_RE = re.compile(r"^\s*(?:miskonsepsi|misconception)\s*[:\-]?\s*(.*)$", re.I)


def parse_paste(text: str) -> list[dict]:
    """Parser deterministik untuk tempelan banyak soal. Tidak ada AI, tidak ada tebakan liar."""
    text = (text or "").replace("\r\n", "\n").strip()
    if not text:
        return []
    blocks = PASTE_SPLIT.split("\n" + text)
    out = []
    for raw in blocks:
        raw = raw.strip()
        if not raw:
            continue
        lines = [l for l in raw.split("\n") if l.strip()]
        stem_lines, options, key, explanation, hints, mis = [], [], "", "", [], None
        for line in lines:
            m = OPT_RE.match(line)
            if m and len(m.group(2).strip()) > 0 and not KEY_RE.match(line):
                options.append(m.group(2).strip())
                continue
            m = KEY_RE.match(line)
            if m:
                key = m.group(1).strip()
                continue
            m = EXP_RE.match(line)
            if m:
                explanation = m.group(1).strip()
                continue
            m = HINT_RE.match(line)
            if m:
                hints.append(m.group(1).strip())
                continue
            m = MIS_RE.match(line)
            if m:
                mis = m.group(1).strip() or None
                continue
            stem_lines.append(re.sub(r"^\s*\d{1,3}\s*[.)]\s*", "", line).strip())
        stem = " ".join(stem_lines).strip()
        if not stem:
            continue
        qtype = "mcq" if len(options) >= 2 else "short_answer"
        if qtype == "mcq" and key:
            k = key.strip()[0].upper()
            key = k if k.isalpha() else key
        out.append({"stem": stem, "options": options, "answer_key": key,
                    "explanation": explanation, "hints": hints, "question_type": qtype,
                    "misconception_id": mis})
    return out


async def _ingest_candidates(rows: list[dict], defaults: dict, actor: str, mode: str) -> dict:
    batch_id = f"BATCH-{uuid.uuid4().hex[:8].upper()}"
    created, rejected = [], []
    for row in rows:
        comp = row.get("competency_id") or defaults.get("competency_id")
        if not comp:
            rejected.append({"row": row, "reason": "competency_id kosong"})
            continue
        try:
            cand = QuestionIn(
                competency_id=comp, stem=row.get("stem", ""),
                question_type=row.get("question_type") or "mcq",
                options=row.get("options") or [],
                answer_key=str(row.get("answer_key") or ""),
                explanation=row.get("explanation") or "",
                hints=row.get("hints") or [],
                difficulty=int(row.get("difficulty") or defaults.get("difficulty") or 2),
                cognitive_level=(row.get("cognitive_level") or defaults.get("cognitive_level") or "C2").upper(),
                misconception_id=row.get("misconception_id"),
                distractor_misconceptions=row.get("distractor_misconceptions") or {},
                estimated_time=int(row.get("estimated_time") or 60),
                is_transfer=bool(row.get("is_transfer")),
                source=mode, status="REVIEW")
            doc = await _dna(cand, actor)
        except HTTPException as e:
            rejected.append({"row": row, "reason": str(e.detail)})
            continue
        doc["issues"] = await validate_doc(doc)
        doc["ingestion_batch_id"] = batch_id
        await db.questions.insert_one(dict(doc))
        doc.pop("_id", None)
        created.append(doc)
    await db.ingestion_batches.insert_one({"id": batch_id, "mode": mode, "actor": actor, "at": now(),
                                           "created": len(created), "rejected": len(rejected),
                                           "question_ids": [c["question_id"] for c in created]})
    return {"batch_id": batch_id, "created": created, "rejected": rejected,
            "note": "Semua hasil ingestion masuk status REVIEW dan wajib ditinjau guru sebelum terbit."}


class PasteIn(BaseModel):
    text: str
    competency_id: str
    difficulty: int = 2
    cognitive_level: str = "C2"


@router.post("/import/paste")
async def import_paste(body: PasteIn, u=Depends(teacher_user)):
    rows = parse_paste(body.text)
    if not rows:
        raise HTTPException(400, "Tidak ada soal yang bisa dibaca dari tempelan ini.")
    return await _ingest_candidates(rows, body.model_dump(), u["user_id"], "paste")


CSV_HEADERS = ["stem", "option_a", "option_b", "option_c", "option_d", "answer_key", "explanation",
               "difficulty", "cognitive_level", "competency_id", "misconception_id", "hint",
               "question_type", "estimated_time", "is_transfer"]


@router.get("/import/template")
async def csv_template():
    return {"headers": CSV_HEADERS,
            "example": ["Hitung 2/3 + 1/6", "5/6", "3/9", "1/2", "2/9", "A",
                        "Samakan penyebut menjadi 6: 4/6 + 1/6 = 5/6", "2", "C2",
                        "KOMP-...", "MIS-PENYEBUT-DIJUMLAH", "Samakan penyebut dulu", "mcq", "60", "false"]}


def _rows_from_table(rows: list[dict]) -> list[dict]:
    out = []
    for r in rows:
        low = { (k or "").strip().lower(): (str(v).strip() if v is not None else "") for k, v in r.items() }
        stem = low.get("stem") or low.get("pertanyaan") or low.get("soal") or ""
        if not stem:
            continue
        options = [low.get(f"option_{c}") or low.get(f"opsi_{c}") or "" for c in "abcde"]
        options = [o for o in options if o]
        out.append({"stem": stem, "options": options,
                    "answer_key": low.get("answer_key") or low.get("kunci") or "",
                    "explanation": low.get("explanation") or low.get("penjelasan") or "",
                    "hints": [h for h in [low.get("hint"), low.get("petunjuk")] if h],
                    "difficulty": low.get("difficulty") or 2,
                    "cognitive_level": low.get("cognitive_level") or "C2",
                    "competency_id": low.get("competency_id") or None,
                    "misconception_id": low.get("misconception_id") or None,
                    "question_type": low.get("question_type") or ("mcq" if options else "short_answer"),
                    "estimated_time": low.get("estimated_time") or 60,
                    "is_transfer": str(low.get("is_transfer") or "").lower() in ("1", "true", "ya", "yes")})
    return out


@router.post("/import/file")
async def import_file(file: UploadFile = File(...), competency_id: str = Form(""),
                      u=Depends(teacher_user)):
    """Bulk import Excel/CSV, plus dokumen PDF/TXT (parser deterministik)."""
    raw = await file.read()
    name = (file.filename or "").lower()
    if name.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        rows = _rows_from_table(list(csv.DictReader(io.StringIO(text))))
        mode = "csv_import"
    elif name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = [str(h or "") for h in next(it)]
        rows = _rows_from_table([dict(zip(header, r)) for r in it])
        mode = "excel_import"
    elif name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
        rows = parse_paste(text)
        mode = "pdf_import"
    elif name.endswith((".txt", ".md")):
        rows = parse_paste(raw.decode("utf-8", errors="replace"))
        mode = "document_import"
    else:
        raise HTTPException(400, "Format tidak didukung. Pakai CSV, XLSX, PDF, atau TXT.")
    if not rows:
        raise HTTPException(400, "Tidak ada soal yang terbaca dari berkas ini.")
    return await _ingest_candidates(rows, {"competency_id": competency_id or None}, u["user_id"], mode)


class GenerateIn(BaseModel):
    competency_id: str
    count: int = 3
    include_transfer: bool = True


@router.post("/generate-candidates")
async def generate_candidates(body: GenerateIn, u=Depends(teacher_user)):
    """Kandidat soal deterministik dari materi + soal terbit yang ada. Selalu masuk REVIEW."""
    comp = await db.curriculum_nodes.find_one({"id": body.competency_id, "type": "competency"}, {"_id": 0})
    if not comp:
        raise HTTPException(404, "kompetensi tidak ditemukan")
    seeds = await db.questions.find({"competency_id": body.competency_id, "is_current": True,
                                      "status": {"$in": ["APPROVED", "PUBLISHED"]}}, {"_id": 0}).to_list(20)
    rows = []
    if seeds:
        for i in range(max(1, min(8, body.count))):
            src = seeds[i % len(seeds)]
            stem = src["stem"]
            if NUM_RE.search(stem):
                stem = NUM_RE.sub(lambda m: str(_shift_num(m.group(0), i + 1)), stem)
                rows.append({"stem": stem, "options": src.get("options"), "answer_key": src.get("answer_key"),
                             "explanation": src.get("explanation"), "difficulty": src["difficulty"],
                             "cognitive_level": src["cognitive_level"], "competency_id": body.competency_id,
                             "question_type": src["question_type"], "hints": src.get("hints")})
            else:
                rows.append({"stem": f"Terapkan pada konteks baru: {stem}", "options": src.get("options"),
                             "answer_key": src.get("answer_key"), "explanation": src.get("explanation"),
                             "difficulty": min(5, src["difficulty"] + 1), "cognitive_level": "C3",
                             "competency_id": body.competency_id, "question_type": src["question_type"],
                             "is_transfer": body.include_transfer, "hints": src.get("hints")})
    else:
        materials = await db.curriculum_nodes.find({"type": "material", "competency_id": body.competency_id},
                                                    {"_id": 0}).to_list(20)
        base = materials[0]["name"] if materials else comp["name"]
        rows.append({"stem": f"Jelaskan dengan kalimatmu sendiri: {base}.",
                     "question_type": "essay", "answer_key": "rubrik",
                     "explanation": f"Kunci: murid menyebut inti {comp['name']}.",
                     "difficulty": 3, "cognitive_level": "C3", "competency_id": body.competency_id})
    return await _ingest_candidates(rows, {"competency_id": body.competency_id}, u["user_id"], "brain_candidate")


@router.get("/review/queue")
async def review_queue(subject_id: str | None = None, limit: int = 100, u=Depends(teacher_user)):
    q: dict[str, Any] = {"is_current": True, "status": {"$in": ["DRAFT", "REVIEW"]}}
    if subject_id:
        q["subject_id"] = subject_id
    items = await db.questions.find(q, {"_id": 0}).sort("created_at", -1).to_list(limit)
    for it in items:
        it["issues"] = await validate_doc(it)
    return {"count": len(items), "items": items}
